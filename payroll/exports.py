"""Finalized payroll export builders.

The export layer deliberately reads the immutable statement snapshot and line
items created for a payroll run. It does not recalculate payroll or infer a
historical employment classification from today's employee profile.
"""

from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO

from .models import PayrollLine, PayrollRun


class PayrollExportError(ValueError):
    """Raised when a finalized run cannot be exported safely."""


MONEY_ZERO = Decimal("0.00")
REGISTER_HEADERS = [
    "Employee code", "Employee", "Pay basis", "Salary rate", "Basic pay",
    "Absent", "Undertime", "Holiday", "Overtime", "Night differential",
    "Allowances", "Other earnings", "Gross pay", "SSS", "PhilHealth",
    "Pag-IBIG", "Withholding tax", "Other deductions", "Total deductions",
    "Net pay", "Employer contributions",
]


def safe_text(value):
    """Prevent spreadsheet and CSV formula injection for text values."""
    text = str(value or "")
    if text[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + text
    return text


def money(value):
    return Decimal(value or "0.00").quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def date_label(value):
    """Portable month/day label for Windows and POSIX Python runtimes."""
    return value.strftime("%b %d, %Y").replace(" 0", " ")


def _sum(lines, predicate):
    return money(sum((line.amount for line in lines if predicate(line)), MONEY_ZERO))


def _salary_rate(snapshot):
    versions = snapshot.get("rate_versions") or snapshot.get("compensation_versions") or {}
    if isinstance(versions, dict):
        for value in versions.values():
            if isinstance(value, dict) and value.get("amount") is not None:
                return money(value.get("amount"))
    return MONEY_ZERO


def _rule_profile_summary(snapshot):
    summaries = []
    seen = set()
    for version in (snapshot.get("rule_versions") or {}).values():
        if not isinstance(version, dict):
            continue
        name = version.get("profile_name") or "Legacy organization rules"
        label = f"{name} (employee override)" if version.get("assignment_id") else f"{name} (inherited)"
        if label not in seen:
            seen.add(label)
            summaries.append(label)
    return ", ".join(summaries) or "No resolved rule profile recorded"


def _reviewed_statutory_amounts(statement):
    """Return reviewed employee and employer amounts by agency.

    New assessments are preferred. The snapshot fallback keeps old finalized
    runs exportable when their assessment rows predate the bulk workspace.
    """
    employee = defaultdict(lambda: MONEY_ZERO)
    employer = defaultdict(lambda: MONEY_ZERO)
    assessments = list(getattr(statement, "statutory_assessments", []).all())
    if assessments and any(item.status != "REVIEWED" for item in assessments):
        raise PayrollExportError(
            f"Statutory review is incomplete for {statement.employee.employee_code}; finalize only after every agency is reviewed."
        )
    for assessment in assessments:
        employee[assessment.agency] += money(assessment.employee_amount)
        employer[assessment.agency] += money(assessment.employer_amount)
    if assessments:
        return employee, employer
    for agency, review in (statement.snapshot or {}).get("statutory_reviews", {}).items():
        if not isinstance(review, dict) or not review.get("reviewed_by_id"):
            continue
        employee[agency] += money((review.get("employee") or {}).get("amount"))
        employer[agency] += money((review.get("employer") or {}).get("amount"))
    return employee, employer


def _statement_row(statement, *, require_classification=True):
    snapshot = statement.snapshot or {}
    employee_snapshot = snapshot.get("employee") or {}
    employment_status = employee_snapshot.get("employment_status")
    if employment_status not in {"REGULAR", "PROBATION"} and require_classification:
        code = employee_snapshot.get("code") or statement.employee.employee_code
        raise PayrollExportError(
            f"Excel export needs a historical Regular or Probation classification for {code}. "
            "Recalculate this payroll with the updated payroll version before finalizing it."
        )
    if employment_status not in {"REGULAR", "PROBATION"}:
        employment_status = "UNCLASSIFIED"

    lines = list(statement.lines.all())
    earnings = [line for line in lines if line.kind == PayrollLine.Kind.EARNING]
    basic = _sum(lines, lambda line: line.code == "REGULAR_PAY")
    overtime = _sum(lines, lambda line: line.code == "OVERTIME_PREMIUM")
    night = _sum(lines, lambda line: line.code == "NIGHT_DIFFERENTIAL")
    holiday = _sum(lines, lambda line: line.code == "DAY_PREMIUM")
    allowances = _sum(lines, lambda line: line.kind == PayrollLine.Kind.EARNING and line.source == "CALCULATED_COMPONENT")
    other_earnings = money(statement.gross_amount) - basic - overtime - night - holiday - allowances
    if other_earnings < MONEY_ZERO:
        raise PayrollExportError(f"The earnings lines for {statement.employee.employee_code} do not reconcile to gross pay.")

    statutory_employee, statutory_employer = _reviewed_statutory_amounts(statement)
    statutory_deductions = sum(statutory_employee.values(), MONEY_ZERO)
    other_deductions = money(statement.deduction_amount) - statutory_deductions
    if other_deductions < MONEY_ZERO:
        raise PayrollExportError(f"The deduction lines for {statement.employee.employee_code} do not reconcile to the statement total.")
    statutory_contributions = sum(statutory_employer.values(), MONEY_ZERO)
    other_contributions = money(statement.employer_contribution_amount) - statutory_contributions
    if other_contributions < MONEY_ZERO:
        raise PayrollExportError(f"The employer contribution lines for {statement.employee.employee_code} do not reconcile to the statement total.")

    employee_code = employee_snapshot.get("code") or statement.employee.employee_code
    employee_name = employee_snapshot.get("name") or statement.employee.full_name
    pay_basis = ""
    for value in (snapshot.get("compensation_versions") or snapshot.get("rate_versions") or {}).values():
        if isinstance(value, dict) and value.get("basis"):
            pay_basis = value["basis"].title()
            break
    row = {
        "group": employment_status,
        "employee_code": safe_text(employee_code),
        "employee": safe_text(employee_name),
        "pay_basis": safe_text(pay_basis or "Hourly"),
        "rule_profiles": safe_text(_rule_profile_summary(snapshot)),
        "salary_rate": _salary_rate(snapshot),
        "basic": basic,
        "absent": MONEY_ZERO,
        "undertime": MONEY_ZERO,
        "holiday": holiday,
        "overtime": overtime,
        "night": night,
        "allowances": allowances,
        "other_earnings": other_earnings,
        "gross": money(statement.gross_amount),
        "sss": statutory_employee["SSS"],
        "philhealth": statutory_employee["PHILHEALTH"],
        "pagibig": statutory_employee["PAGIBIG"],
        "withholding": statutory_employee["WITHHOLDING"],
        "other_deductions": other_deductions,
        "deductions": money(statement.deduction_amount),
        "net": money(statement.net_amount),
        "employer_contributions": money(statement.employer_contribution_amount),
        "employer_sss": statutory_employer["SSS"],
        "employer_philhealth": statutory_employer["PHILHEALTH"],
        "employer_pagibig": statutory_employer["PAGIBIG"],
        "other_contributions": other_contributions,
        "lines": [
            {"kind": line.get_kind_display(), "label": safe_text(line.label), "amount": money(line.amount)}
            for line in lines
        ],
    }
    if money(row["gross"] - row["deductions"]) != row["net"]:
        raise PayrollExportError(f"The gross, deductions, and net pay do not reconcile for {employee_code}.")
    return row


def build_finalized_export_data(run, *, strict_classification=True):
    if run.status != PayrollRun.Status.FINALIZED:
        raise PayrollExportError("Only finalized payroll can be exported.")
    statements = list(run.statements.select_related("employee").prefetch_related(
        "lines", "statutory_assessments"
    ).order_by("employee__last_name", "employee__first_name", "pk"))
    if strict_classification:
        unsupported = []
        for statement in statements:
            snapshot_status = (statement.snapshot or {}).get("employee", {}).get("employment_status")
            if snapshot_status not in {"REGULAR", "PROBATION"}:
                code = (statement.snapshot or {}).get("employee", {}).get("code") or statement.employee.employee_code
                unsupported.append(f"{code} ({snapshot_status or 'missing classification'})")
        if unsupported:
            raise PayrollExportError(
                "Excel export needs a historical Regular or Probation classification for "
                f"{len(unsupported)} employee(s): {', '.join(unsupported)}. "
                "Recalculate a draft with the updated payroll version before finalizing it."
            )
    rows = [_statement_row(statement, require_classification=strict_classification) for statement in statements]
    groups = {"REGULAR": [], "PROBATION": []}
    for row in rows:
        if row["group"] in groups:
            groups[row["group"]].append(row)
    return {
        "run": run,
        "rows": rows,
        "groups": groups,
        "period": f"{date_label(run.period_start)} – {date_label(run.period_end)}",
    }


def csv_export_rows(data):
    rows = []
    for row in data["rows"]:
        rows.append([
            row["employee_code"], row["employee"], row["rule_profiles"], data["run"].period_start.isoformat(),
            data["run"].period_end.isoformat(), data["run"].pay_date.isoformat(),
            data["run"].get_pay_frequency_display(), data["run"].currency,
            f"{row['basic']:.2f}", f"{row['overtime']:.2f}", f"{row['night']:.2f}",
            f"{row['holiday']:.2f}", f"{row['allowances']:.2f}", f"{row['other_earnings']:.2f}",
            f"{row['deductions']:.2f}", f"{row['employer_contributions']:.2f}",
            f"{row['gross']:.2f}", f"{row['net']:.2f}",
        ])
    return rows


def _write_register(ws, data, title, rows):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    navy = "15243D"
    blue = "EAF2FF"
    border = Border(bottom=Side(style="thin", color="D8E1EF"))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(REGISTER_HEADERS))
    ws.cell(1, 1, f"{data['run'].organization.name} · {title} payroll register").font = Font(size=16, bold=True, color=navy)
    ws.cell(2, 1, f"{data['run'].reference} · {data['period']} · Pay date {date_label(data['run'].pay_date)}").font = Font(color="526B8A")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(REGISTER_HEADERS))
    ws.cell(4, 1, "Employee payroll register").font = Font(size=12, bold=True, color=navy)
    header_row = 6
    for col, header in enumerate(REGISTER_HEADERS, 1):
        cell = ws.cell(header_row, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2463EB")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for index, row in enumerate(rows, header_row + 1):
        values = [row["employee_code"], row["employee"], row["pay_basis"], row["salary_rate"], row["basic"], row["absent"], row["undertime"], row["holiday"], row["overtime"], row["night"], row["allowances"], row["other_earnings"], row["gross"], row["sss"], row["philhealth"], row["pagibig"], row["withholding"], row["other_deductions"], row["deductions"], row["net"], row["employer_contributions"]]
        for col, value in enumerate(values, 1):
            cell = ws.cell(index, col, value if not isinstance(value, str) else safe_text(value))
            cell.border = border
            if col >= 4:
                cell.number_format = '#,##0.00'
    total_row = header_row + len(rows) + 1
    ws.cell(total_row, 1, "TOTAL").font = Font(bold=True, color=navy)
    for col in range(4, len(REGISTER_HEADERS) + 1):
        total = sum((row["salary_rate"] if col == 4 else row[{5: "basic", 6: "absent", 7: "undertime", 8: "holiday", 9: "overtime", 10: "night", 11: "allowances", 12: "other_earnings", 13: "gross", 14: "sss", 15: "philhealth", 16: "pagibig", 17: "withholding", 18: "other_deductions", 19: "deductions", 20: "net", 21: "employer_contributions"}[col]] for row in rows), MONEY_ZERO)
        ws.cell(total_row, col, total).font = Font(bold=True, color=navy)
        ws.cell(total_row, col).number_format = '#,##0.00'
    ws.freeze_panes = "A7"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(REGISTER_HEADERS))}{header_row + max(len(rows), 1)}"
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = "1:6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    widths = [16, 26, 12, 14] + [14] * (len(REGISTER_HEADERS) - 4)
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _write_payslips(ws, data, title, rows):
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.styles import Alignment, Font, PatternFill, Side, Border

    navy = "15243D"
    row_number = 1
    for index, row in enumerate(rows):
        ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=4)
        ws.cell(row_number, 1, f"{data['run'].organization.name} · {title} payslip").font = Font(size=16, bold=True, color=navy)
        row_number += 1
        ws.cell(row_number, 1, safe_text(row["employee"])).font = Font(size=13, bold=True, color=navy)
        ws.cell(row_number, 3, f"#{row['employee_code']}").font = Font(color="526B8A")
        row_number += 1
        ws.cell(row_number, 1, f"{data['run'].reference} · {data['period']}").font = Font(color="526B8A")
        ws.cell(row_number, 3, f"Pay date {date_label(data['run'].pay_date)}").font = Font(color="526B8A")
        row_number += 2
        for label, amount in (("Earnings", None), ("Basic pay", row["basic"]), ("Overtime", row["overtime"]), ("Night differential", row["night"]), ("Holiday", row["holiday"]), ("Allowances", row["allowances"]), ("Other earnings", row["other_earnings"]), ("Gross pay", row["gross"]), ("Deductions", None), ("SSS", row["sss"]), ("PhilHealth", row["philhealth"]), ("Pag-IBIG", row["pagibig"]), ("Withholding tax", row["withholding"]), ("Other deductions", row["other_deductions"]), ("Total deductions", row["deductions"]), ("Net pay", row["net"]), ("Employer contributions", row["employer_contributions"])):
            ws.cell(row_number, 1, label)
            if amount is None:
                ws.cell(row_number, 1).font = Font(bold=True, color="FFFFFF")
                for col in range(1, 5):
                    ws.cell(row_number, col).fill = PatternFill("solid", fgColor="2463EB")
            else:
                ws.cell(row_number, 1).font = Font(color=navy)
                ws.cell(row_number, 4, amount).number_format = '#,##0.00'
                ws.cell(row_number, 4).font = Font(bold=label in {"Gross pay", "Total deductions", "Net pay"}, color=navy)
            row_number += 1
        row_number += 2
        ws.cell(row_number, 1, "Employer review signature").font = Font(color="526B8A")
        row_number += 2
        if index < len(rows) - 1:
            ws.row_breaks.append(Break(id=row_number))
        row_number += 2
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 3
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 18
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1


def build_finalized_workbook(data):
    """Return an in-memory workbook with exactly the four requested sheets."""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, group in (("Regular", "REGULAR"), ("Probation", "PROBATION")):
        sheet = workbook.create_sheet(title)
        _write_register(sheet, data, title, data["groups"][group])
    for title, group in (("Payslip Regular", "REGULAR"), ("Payslip Probation", "PROBATION")):
        sheet = workbook.create_sheet(title)
        _write_payslips(sheet, data, title.replace("Payslip ", ""), data["groups"][group])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
